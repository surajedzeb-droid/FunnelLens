"""Excel workbook writing (README section 9): one sheet per selected report (in the given
order), a Raw Data sheet, a Run Info sheet, and an Issues sheet if checks.py found any.

write_workbook's `results` list must include the "raw" ReportResult (from
funnellens.reports.raw) alongside whatever reports the user selected -- it is pulled out
and placed as its own Raw Data sheet rather than counted among the selected reports.

`run_info` is a plain dict the caller builds (Phase 6/7 know things export.py can't: which
filters/presets were used, the API call count, ...). Two entries drive filename/formatting
specifically: `from_date` and `to_date` (date objects); everything else in the dict is just
rendered as a Run Info row. Stage-basis labels and checks.py issues are always computed here
from `results`, not expected from the caller.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from funnellens.checks import run_checks
from funnellens.reports import ReportResult
from funnellens.timeutil import IST

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "output"

_HEADER_FONT = Font(bold=True)
_TOTAL_FONT = Font(bold=True)
_TOTAL_FILL = PatternFill(start_color="FFEFEFEF", end_color="FFEFEFEF", fill_type="solid")
_PCT_FORMAT = '0.00"%"'  # values are already 0-100 (e.g. safe_pct's 66.7), not a 0-1 fraction
_INT_FORMAT = "0"
_MAX_COLUMN_WIDTH = 40


def _is_pct_column(name: object) -> bool:
    return isinstance(name, str) and name.strip().endswith("%")


def _sheet_name(title: str, used: set[str]) -> str:
    name = title[:31]
    suffix = 1
    while name in used:
        suffix += 1
        name = f"{title[:29]}_{suffix}"[:31]
    used.add(name)
    return name


def _write_report_sheet(wb: Workbook, result: ReportResult) -> None:
    ws = wb.create_sheet(_sheet_name(result.title, {s.title for s in wb.worksheets}))
    df = result.dataframe
    columns = [c for c in df.columns if c != "is_total"]

    ws.cell(row=1, column=1, value=result.title).font = Font(bold=True, size=14)
    header_row = 2
    if result.stage_basis_label:
        ws.cell(row=2, column=1, value=result.stage_basis_label).font = Font(italic=True)
        header_row = 3

    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=name)
        cell.font = _HEADER_FONT
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    is_total = df["is_total"] if "is_total" in df.columns else pd.Series(False, index=df.index)
    for row_offset, (_, record) in enumerate(df.iterrows()):
        row_idx = header_row + 1 + row_offset
        total = bool(is_total.iloc[row_offset])
        for col_idx, name in enumerate(columns, start=1):
            value = record[name]
            if pd.isna(value):
                value = ""
            elif getattr(value, "tzinfo", None) is not None:
                value = value.replace(tzinfo=None)  # Excel doesn't support tz-aware datetimes
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if _is_pct_column(name) and isinstance(value, (int, float)):
                cell.number_format = _PCT_FORMAT
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = _INT_FORMAT
            if total:
                cell.font = _TOTAL_FONT
                cell.fill = _TOTAL_FILL

    _autosize_columns(ws, columns, header_row)


def _autosize_columns(ws: Worksheet, columns: list, header_row: int) -> None:
    for col_idx, name in enumerate(columns, start=1):
        width = len(str(name))
        for cell in ws.iter_rows(min_row=header_row + 1, min_col=col_idx, max_col=col_idx):
            width = max(width, len(str(cell[0].value)) if cell[0].value is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, _MAX_COLUMN_WIDTH)


def _write_run_info_sheet(wb: Workbook, run_info: dict, stage_labels: dict[str, str], issue_count: int) -> None:
    ws = wb.create_sheet(_sheet_name("Run Info", {s.title for s in wb.worksheets}))
    ws.cell(row=1, column=1, value="Run Info").font = Font(bold=True, size=14)

    rows = list(run_info.items())
    if stage_labels:
        rows.append(("Stage-basis labels", "; ".join(f"{title}: {label}" for title, label in stage_labels.items())))
    rows.append(("Issues found", issue_count))

    row_idx = 2
    for key, value in rows:
        ws.cell(row=row_idx, column=1, value=str(key)).font = _HEADER_FONT
        if isinstance(value, dict):
            value = ", ".join(f"{k}={v}" for k, v in value.items())
        elif isinstance(value, (list, tuple, set)):
            value = ", ".join(str(v) for v in value)
        elif getattr(value, "tzinfo", None) is not None:
            # README section 9: generated-at time is recorded in IST. Excel also rejects
            # tz-aware datetimes outright, so convert then isoformat rather than just stripping.
            value = value.astimezone(IST).isoformat()
        ws.cell(row=row_idx, column=2, value=value)
        row_idx += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80


def _write_issues_sheet(wb: Workbook, issues: list[str]) -> None:
    ws = wb.create_sheet(_sheet_name("Issues", {s.title for s in wb.worksheets}))
    ws.cell(row=1, column=1, value="Issue").font = _HEADER_FONT
    ws.freeze_panes = "A2"
    for row_idx, issue in enumerate(issues, start=2):
        ws.cell(row=row_idx, column=1, value=issue)
    ws.column_dimensions["A"].width = min(max((len(i) for i in issues), default=20) + 2, 120)


def write_workbook(results: list[ReportResult], run_info: dict, path: str | Path | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    raw_result = next((r for r in results if r.key == "raw"), None)
    selected = [r for r in results if r.key != "raw"]

    for result in selected:
        _write_report_sheet(wb, result)
    if raw_result is not None:
        _write_report_sheet(wb, raw_result)

    stage_labels = {r.title: r.stage_basis_label for r in results if r.stage_basis_label}
    issues = run_checks({r.key: r for r in results})
    _write_run_info_sheet(wb, run_info, stage_labels, len(issues))
    if issues:
        _write_issues_sheet(wb, issues)

    buffer = BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()

    if path is None:
        from_date, to_date = run_info.get("from_date"), run_info.get("to_date")
        path = OUTPUT_DIR / f"FunnelLens_{from_date}_to_{to_date}.xlsx"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)

    return data
