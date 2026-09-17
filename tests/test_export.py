from datetime import date, datetime, timezone

import openpyxl
import pandas as pd

from funnellens.export import write_workbook
from funnellens.reports import ReportContext
from funnellens.reports.lead_funnel import build as build_lead_funnel
from funnellens.reports.raw import build as build_raw
from funnellens.reports.stage_wise import build as build_stage

from tests.fixtures.reports_fixture import DAY1, DAY2, RUN_AT, build_fixture_dataset
from tests.test_reports import make_settings


def make_results():
    dataset = build_fixture_dataset()
    context = ReportContext(settings=make_settings(), run_at=RUN_AT)
    return [
        build_lead_funnel(dataset, DAY1, DAY2, context),
        build_stage(dataset, DAY1, DAY2, context),
        build_raw(dataset, DAY1, DAY2, context),
    ]


def make_run_info():
    return {"from_date": DAY1, "to_date": DAY2, "generated_at": RUN_AT, "reports": ["lead_funnel", "stage"],
            "excluded_owners_included": False, "api_calls": 42}


def test_write_workbook_returns_bytes():
    data = write_workbook(make_results(), make_run_info())
    assert isinstance(data, bytes)
    assert data[:2] == b"PK"  # xlsx is a zip archive


def test_write_workbook_creates_expected_sheets(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Lead Funnel", "Stage Wise", "Raw Data", "Run Info"]


def test_total_rows_are_bold_with_fill(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Lead Funnel"]
    total_cell = next(c for row in ws.iter_rows(min_row=3) for c in row if c.value == "Total")
    assert total_cell.font.bold is True
    assert total_cell.fill.start_color.rgb == "00FFEFEF" or total_cell.fill.start_color.rgb == "FFEFEFEF"


def test_percentage_columns_use_percent_format(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Lead Funnel"]
    header_row = [c.value for c in ws[2]]
    pct_col = header_row.index("Overdues %") + 1
    data_cell = ws.cell(row=4, column=pct_col)
    assert "%" in data_cell.number_format


def test_header_row_is_frozen(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    wb = openpyxl.load_workbook(path)
    assert wb["Lead Funnel"].freeze_panes == "A3"  # title row 1, header row 2


def test_stage_wise_sheet_has_stage_as_of_label(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Stage Wise"]
    assert ws.cell(row=2, column=1).value == "Stage as of 2026-09-03"
    assert ws.freeze_panes == "A4"  # title + label + header


def test_run_info_sheet_contains_given_fields(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Run Info"]
    keys = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert "api_calls" in keys
    assert "Stage-basis labels" in keys
    assert "Issues found" in keys


def test_no_issues_means_no_issues_sheet(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    wb = openpyxl.load_workbook(path)
    assert "Issues" not in wb.sheetnames


def test_file_is_saved_to_the_given_path(tmp_path):
    path = tmp_path / "custom.xlsx"
    write_workbook(make_results(), make_run_info(), path=path)
    assert path.exists()


def test_default_path_follows_readme_pattern(monkeypatch, tmp_path):
    import funnellens.export as export_module
    monkeypatch.setattr(export_module, "OUTPUT_DIR", tmp_path)
    write_workbook(make_results(), make_run_info())
    expected = tmp_path / f"FunnelLens_{DAY1}_to_{DAY2}.xlsx"
    assert expected.exists()


def test_timezone_aware_datetime_cells_do_not_crash_the_writer(tmp_path):
    raw = make_results()[2]  # the raw ReportResult
    raw.dataframe["enrolled_date"] = pd.Series(
        [datetime(2026, 9, 1, 5, 0, tzinfo=timezone.utc)] * len(raw.dataframe))
    path = tmp_path / "out.xlsx"
    write_workbook([raw], make_run_info(), path=path)  # must not raise
    wb = openpyxl.load_workbook(path)
    assert wb["Raw Data"].cell(row=3, column=1).value is not None


def test_sheet_names_are_truncated_to_31_chars():
    from funnellens.export import _sheet_name
    name = _sheet_name("A Very Long Report Title That Exceeds The Excel Limit", set())
    assert len(name) <= 31
